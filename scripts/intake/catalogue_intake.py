#!/usr/bin/env python3
"""Supplier WhatsApp chat exports -> review sheet -> catalogue.  See docs/CATALOGUE_INTAKE.md.

  scan      find new "WhatsApp Chat - <supplier>.zip" exports, parse the posts and write a review
            spreadsheet (intake/review/*.xlsx). Needs no secrets and never touches the live shop.
  backfill  one-off: add sourcing records for products that predate this tool.
  contact   make labelled contact sheets of photo-only posts so they can be named by eye.
  setnames  fill product names into a review sheet (used for photo-only posts).
  stock     stock check only, over every export (already-scanned ones too): writes the stock-check sheet.
  log       show the run log (intake/RUN_LOG.md): when each scan/apply ran and what it did.
  apply     push rows you marked ADD to the shop, and record new supplier costs for rows marked
            UPDATE COST (selling prices are NOT changed here: admin sets them). Needs SUPABASE_SERVICE_KEY in the
            environment for that one run; the key is never stored anywhere.

Third-party chats and photos stay in intake/ (git-ignored) and are never published.
"""
import argparse, glob, hashlib, json, os, re, shutil, sys, time, uuid, zipfile, datetime, collections
import urllib.request, urllib.error
from PIL import Image, ImageFile
ImageFile.LOAD_TRUNCATED_IMAGES = True
import openpyxl
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, "..", ".."))
CFG = json.load(open(os.path.join(HERE, "config.json")))

# ---------------------------------------------------------------- site knowledge
CATEGORIES = {  # slug -> keywords that suggest it (first match wins, order matters)
    "sneakers-shoes": r"sneaker|puma|nike|adidas|vans|jordan|\bboots?\b|dunk|air ?max|air ?force|shox|timberland|clarks|wallabee|lugz|new balance|dr martens|samba|campus|sandals?|casuals?|brogues?|loafers?|trainers?|\bshoes?\b|mules|slides|crocs",
    "motorcycle-gear": r"helmet|knee guard|elbow guard|riding|gloves|reflective|rain suit|motocross|motorcycle|biker",
    "audio-speakers": r"speaker|soundbar|woofer|home theat|audio system|headphone|earbud|microphone|\bmic\b",
    "home-appliances": r"microwave|washing machine|twin tub|cooker|stove|\bfan\b|heater|vacuum|steam cleaner|kettle|blender|roti|fridge|refrigerator|\boven\b|hot ?plate|iron box|air ?fryer|dispenser",
    "personal-care": r"shaver|clipper|trimmer|blow ?dryer|hair dryer|straightener|perfume|body mist|gift set",
    "kitchen-dining": r"cutlery|cookware|\bpots?\b|\bpans?\b|lunch box|mug|plates?|glasses|jar|flask|bottle|dinner set|sugar dish",
    "interior-design": r"curtain|blanket|duvet|bed ?sheet|pillow|flower|vase|carpet|rug|wall (?:art|decor)|mirror",
    "home-storage": r"\brack\b|organi[sz]er|hanger|shoe slot|storage|basket|luggage|suitcase",
    "tools-car-care": r"drill|tool box|ladder|\btape\b|lock box|floodlight|solar|car wash|foam cleaner|screwdriver",
    "toy-shop": r"\btoys?\b|\bdolls?\b|teddy|puzzle|magic cube|twist cube|lego|building blocks|bubble gun|toy car|remote control car|stroller|play ?set|play mat|kids kitchen|money safe|fairy wings|birthday (?:crown|banner)|balloons?",
    "sports-leisure": r"skateboard|skates|camping|umbrella|football|gym|dumbbell|yoga",
    "furniture": r"\btable\b|\bchair\b|sofa|\bbed\b|wardrobe|shelf|desk",
    "electronics": r"watch|router|wifi|calculator|extension|tripod|tv (?:bracket|mount)|bracket|phone|power ?bank|charger|cable|camera|smart|bluetooth|tablet|laptop",
}
DESC = {
    "sneakers-shoes": "Comfortable everyday footwear.", "audio-speakers": "Sound that fills the room.",
    "home-appliances": "A dependable appliance for your home.", "electronics": "A practical gadget for home, office or on the go.",
    "tools-car-care": "A handy tool for the job.", "personal-care": "Personal grooming made easy.",
    "motorcycle-gear": "Riding gear built for the road.", "kitchen-dining": "For your kitchen and table.",
    "home-storage": "Keeps your space tidy and organised.", "interior-design": "Finishing touches for your home.",
    "sports-leisure": "For play and the outdoors.", "toy-shop": "Toys and party fun for kids.", "furniture": "Furniture for your home or office.",
}
BIG_BRANDS = re.compile(r"(?i)\b(nike|adidas|puma|new balance|vans|jordan|timberland|clarks|dr\.? ?martens|lugz|cat|converse|skechers|reebok|under armour|on cloud|air ?max|air ?force)\b")
GREETING = re.compile(r"(?i)\b(good morning|happy (?:monday|tuesday|wednesday|thursday|friday|saturday|sunday)|blessings?|prayer|amen|father|lord|heavenly|almighty|god|happy new)\b")
SOLD_OUT = re.compile(r"(?i)\b(sold ?out|out of stock|out-of-stock|no stock|imeisha|zimeisha|stock (?:is )?(?:over|finished)|not available)\b")
NEW_ITEM_DAYS = CFG.get("new_item_days", 7)        # only posts this recent are offered as new products
STOCK_CHECK_DAYS = CFG.get("stock_check_days", 21)  # on-site products not reposted this long get a stock warning
STALE_EXPORT_DAYS = 14                              # warn when a supplier's chat hasn't been exported for this long
EMO = re.compile(r"[\U00010000-\U0010ffff☀-➿⬀-⯿️‍]")


def cfg_path(p):
    return os.path.expanduser(p)


LOG_HEAD = "# Catalogue intake run log\n\nNewest at the bottom. Written by scripts/intake/catalogue_intake.py.\n\n| When | Command | Result |\n|---|---|---|\n"


def log_run(root, cmd, result):
    """One line per run in intake/RUN_LOG.md, so there's always a record of the last update."""
    path = os.path.join(root, "RUN_LOG.md")
    new = not os.path.exists(path)
    with open(path, "a", encoding="utf-8") as f:
        if new: f.write(LOG_HEAD)
        when = datetime.datetime.now().strftime("%Y-%m-%d %H:%M")
        f.write(f"| {when} | {cmd} | {result.replace('|', '/')} |\n")


def cmd_log(args):
    path = os.path.join(work_root(args), "RUN_LOG.md")
    if not os.path.exists(path): print("No runs logged yet."); return
    lines = [l for l in open(path, encoding="utf-8").read().splitlines() if l.startswith("| 2")]
    print("\n".join(lines[-args.last:]) if lines else "No runs logged yet.")
    changed = lambda l: "| apply |" in l and "dry run" not in l and bool(re.search(r"added [1-9]|updated [1-9]", l))
    applied = [l for l in lines if changed(l)]
    print("\nLast catalogue change:", applied[-1].split("|")[1].strip() if applied else "none logged yet")
    spath = os.path.join(work_root(args), "stock_seen.json")
    if os.path.exists(spath):
        sups = json.load(open(spath)).get("suppliers", {})
        if sups:
            print("Newest message per supplier chat:", "; ".join(f"{k} {v.get('latest_chat', '?')}" for k, v in sorted(sups.items())))


def work_root(args):
    return os.path.abspath(args.work) if getattr(args, "work", None) else os.path.join(ROOT, CFG["work_dir"])


def load_env():
    env = {}
    p = os.path.join(ROOT, ".env.local")
    if os.path.exists(p):
        for line in open(p):
            if "=" in line and not line.startswith("#"):
                k, v = line.strip().split("=", 1); env[k] = v
    return env


def http(method, url, headers, data=None, tries=4):
    for i in range(tries):
        try:
            r = urllib.request.Request(url, data=data, headers=headers, method=method)
            with urllib.request.urlopen(r, timeout=60) as resp:
                return resp.status, resp.read()
        except urllib.error.HTTPError as e:
            return e.code, e.read()
        except Exception:
            if i == tries - 1: raise
            time.sleep(2 * (i + 1))


def live_products():
    env = load_env()
    base, key = env.get("NEXT_PUBLIC_SUPABASE_URL"), env.get("NEXT_PUBLIC_SUPABASE_ANON_KEY")
    if not base or not key:
        return []
    st, body = http("GET", f"{base}/rest/v1/products?select=id,slug,name,price,category", {"apikey": key, "Authorization": "Bearer " + key})
    return json.loads(body) if st == 200 else []


# ---------------------------------------------------------------- parsing
LINE = [  # iOS:  [25/07/2026, 16:53:12] Name: text      Android:  25/07/2026, 16:53 - Name: text
    re.compile(r"^\[(\d{1,2})/(\d{1,2})/(\d{2,4}), (\d{1,2}):(\d{2})(?::\d{2})?\s?(?:[AP]M)?\] ([^:]+): (.*)$", re.S),
    re.compile(r"^(\d{1,2})/(\d{1,2})/(\d{2,4}), (\d{1,2}):(\d{2})(?:\s?[AP]M)? - ([^:]+): (.*)$", re.S),
]
ATTACH = re.compile(r"<attached: ([^>]+)>|^(IMG-[\w-]+\.(?:jpg|jpeg|png|webp)) \(file attached\)", re.M)
START = re.compile(r"(?m)^(?=\[?\d{1,2}/\d{1,2}/\d{2,4}, \d{1,2}:\d{2})")


def parse_chat(text):
    text = text.replace("‎", "").replace(" ", " ")
    msgs = []
    for block in START.split(text):
        block = block.strip("\n")
        for rx in LINE:
            m = rx.match(block)
            if m:
                d, mo, y, H, M, sender, body = m.groups()
                y = int(y) + (2000 if int(y) < 100 else 0)
                try:
                    ts = datetime.datetime(y, int(mo), int(d), int(H), int(M))
                except ValueError:
                    break
                msgs.append({"ts": ts, "sender": sender.strip("⁦⁩‪‬ ").strip(), "body": body})
                break
    return msgs


def clean_text(b):
    return re.sub(r"\s+", " ", ATTACH.sub("", b)).strip()


PRICE_PATTERNS = re.compile(
    r"(?:@|=|ksh\.?s?|kes|price|only|now|selling)\s*[:\-]?\s*(\d{1,3}(?:,\d{3})+|\d{3,6})(?:\s*/=|/-)?"
    r"|(\d{1,3}(?:,\d{3})+|\d{3,6})\s*(?:/=|/-|ksh|bob)|(\d{1,3})\s?k\b|(?:^|\s)(\d{4,5})\s*$", re.I)


def prices_in(text):
    t = text.replace("*", "").replace("_", "")
    out = []
    for m in PRICE_PATTERNS.finditer(t):
        g = [x for x in m.groups() if x]
        if not g: continue
        v = int(g[0].replace(",", ""))
        if m.group(3): v *= 1000
        if 50 <= v <= 2_000_000: out.append(v)
    return out


def clean_name(t):
    t = EMO.sub("", t).replace("*", "").replace("_", "").replace("~", " ")
    t = re.sub(r"(?i)\b(restocked|new in|new arrivals?|back in stock|now available|available|price drop|offer price:?|wholesale price|quality|hot deal|limited stock|original|genuine|authentic)\b", "", t)
    t = re.sub(r"(?i)\bsizes?\s*[:\-]?\s*[\d, \-]+", "", t)
    t = re.sub(r"(?i)(@|=|ksh\.?s?|price|only|now|selling)\s*[:\-]?\s*[\d,\.]+\s*(/=|/-|k)?", "", t)
    t = re.sub(r"\d{9,12}", "", t)
    t = re.sub(r"\s+", " ", t).strip(" -:.,;!?@")
    t = re.split(r"\s[-–—•▪◼✅🔹💥]\s|[•▪◼✅]|\.\s", t)[0]
    t = re.sub(r"\s+(Comes with|Removable|Water and wind|Fully padded|Powerful|Weight|featuring|Selling|Durable|Code:)\b.*$", "", t, flags=re.I)
    return " ".join(t.split()[:9]).strip(" -:.,;!?@")


def sizes_in(t):
    m = re.search(r"(?i)sizes?\s*[:\-]?\s*(\d{2}(?:\s*[-,]\s*\d{2})*)", t)
    return re.sub(r"\s", "", m.group(1)) if m else ""


def guess_category(text):
    for slug, rx in CATEGORIES.items():
        if re.search(rx, text, re.I): return slug
    return ""


def build_items(msgs, photos_dir, supplier):
    items, last_text = [], None
    for m in sorted(msgs, key=lambda x: x["ts"]):
        if m["sender"] in CFG["skip_senders"]: continue
        photos = [a or b for a, b in ATTACH.findall(m["body"])]
        txt = clean_text(m["body"])
        if "omitted" in txt and not photos:
            continue
        if photos:
            use = txt or (last_text[1] if last_text and (m["ts"] - last_text[0]).total_seconds() <= 600 else "")
            files = [p for p in photos if os.path.exists(os.path.join(photos_dir, p)) and p.lower().endswith((".jpg", ".jpeg", ".png", ".webp"))]
            if files:
                items.append({"ts": m["ts"], "text": use, "own": bool(txt), "photos": files, "sender": m["sender"]})
        elif txt and not GREETING.search(txt):
            last_text = (m["ts"], txt)
    merged = []
    for it in items:  # photos of one product posted back-to-back
        if merged and not it["own"] and (it["ts"] - merged[-1]["ts"]).total_seconds() <= 180 and (not it["text"] or it["text"] == merged[-1]["text"]):
            merged[-1]["photos"] += it["photos"]; continue
        merged.append(it)
    out = []
    for it in merged:
        if GREETING.search(it["text"]) and not prices_in(it["text"]): continue  # prayers, good-morning posters
        pr = prices_in(it["text"])
        name = clean_name(it["text"])
        out.append({"ts": it["ts"], "supplier": supplier, "text": it["text"], "name": name, "prices": pr,
                    "price": pr[0] if pr else None, "sizes": sizes_in(it["text"]), "photos": it["photos"],
                    "category": guess_category(it["text"] or name)})
    return out


def canonical_supplier(group):
    """(display name, kind): kind is 'supplier' (cost + margin), 'retail' (your own group) or 'unlisted'."""
    low = group.lower()
    for k in CFG["retail_price_sources"]:
        if k.lower() in low: return CFG["own_group_label"], "retail"
    for s in CFG["suppliers"]:
        if any(m in low for m in s["match"]): return s["name"], "supplier"
    return group, "unlisted"


def sell_price(cost, kind="supplier"):
    if not cost: return None
    # Your own broadcast groups already post retail prices, so no margin on those.
    p = cost * (1 + (0 if kind == "retail" else CFG["markup_percent"]) / 100.0)
    step = CFG["round_up_to"]
    return max(CFG["min_sell_price"], int(-(-p // step) * step))


def cost_check(cost, site_price, kind):
    """A supplier reposted something you already sell. The scan has no access to the recorded
    (admin-only) cost, so it compares the newly posted cost with your current selling price."""
    if kind == "retail" or not cost or not site_price:
        return "SKIP (on site)"
    margin = (site_price - cost) / cost * 100
    target = sell_price(cost, kind)
    if site_price < cost:
        return f"⚠ BELOW COST: cost {cost:,} · your price {site_price:,} · suggested {target:,}"
    if margin < CFG["markup_percent"] / 2:
        return f"⚠ THIN MARGIN: cost {cost:,} · your price {site_price:,} · margin {margin:.0f}% · suggested {target:,}"
    if site_price > target * 1.15:
        return f"↓ ROOM TO CUT: cost {cost:,} · your price {site_price:,} · margin {margin:.0f}% · suggested {target:,}"
    return "SKIP (on site)"


def tokens(s):
    stop = {"the", "and", "with", "for", "size", "pcs", "pc", "set", "kg", "black", "white", "new"}
    return {t for t in re.findall(r"[a-z0-9]+", s.lower()) if len(t) > 1 and t not in stop}


def match_live(name, live):
    a = tokens(name)
    if len(a) < 2: return None
    best = (0, None)
    for p in live:
        b = tokens(p["name"])
        if not b: continue
        j = len(a & b) / min(len(a), len(b))
        if j > best[0]: best = (j, p)
    return best[1] if best[0] >= 0.66 else None


# Only the chat text and photos are ever used, so videos, voice notes and documents (the bulk of a
# big export) are left in the zip instead of filling the disk.
EXTRACT_EXT = (".txt", ".jpg", ".jpeg", ".png", ".webp")


def file_sha1(path):
    """Hash in chunks: a 13 GB export must never be read into memory in one go."""
    h = hashlib.sha1()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def fully_local(path):
    """Google Drive can leave a big export as a partly downloaded placeholder; reading one blocks
    forever. A file whose on-disk size is well below its real size is not ready yet."""
    st = os.stat(path)
    return st.st_blocks * 512 >= st.st_size * 0.98


def safe_extract(zpath, dest):
    os.makedirs(dest, exist_ok=True)
    root = os.path.realpath(dest)
    with zipfile.ZipFile(zpath) as z:
        for info in z.infolist():
            target = os.path.realpath(os.path.join(dest, info.filename))
            if not target.startswith(root + os.sep) and target != root:
                raise SystemExit(f"unsafe path in zip: {info.filename}")
        for info in z.infolist():
            if info.filename.lower().endswith(EXTRACT_EXT):
                z.extract(info, dest)


def slug(t):
    return re.sub(r"[^a-z0-9]+", "-", t.lower()).strip("-")[:60] or "item"


# ---------------------------------------------------------------- scan
HEAD = ["Decision", "Photo", "Item ID", "Supplier", "Posted", "Product name", "Category", "Supplier price (cost)",
        "Proposed sell price", "Final sell price", "Sizes", "Suggested", "Already on site", "Site price",
        "Source group", "Photo files", "Original message", "Applied", "Site product ID"]
COL = {h: i + 1 for i, h in enumerate(HEAD)}


def write_sheet(path, rows):
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Review"
    ws.append(HEAD)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="4D66B0"); c.alignment = Alignment(wrap_text=True, vertical="center")
    for r in rows:
        ws.append([r["decision"], None, r["id"], r["supplier"], r["posted"], r["name"], r["category"], r["cost"], r["proposed"], r["final"],
                   r["sizes"], r["suggested"], r["site_name"], r["site_price"], r["group"], ", ".join(r["files"]), r["text"][:400], "", r["site_id"]])
        row = ws.max_row
        ws.row_dimensions[row].height = 70
        try:
            im = XLImage(r["thumb"]); im.width = im.height = 88
            ws.add_image(im, f"B{row}")
        except Exception:
            pass
        if r["suggested"].startswith(("⚠", "↓")):
            for c in ws[row]: c.fill = PatternFill("solid", fgColor="FFE0B2" if r["suggested"].startswith("⚠") else "E3F2FD")
        elif r["suggested"].startswith("SKIP"):
            for c in ws[row]: c.fill = PatternFill("solid", fgColor="EEEEEE")
        elif not r["name"] or r["name"].startswith("("):
            ws.cell(row=row, column=COL["Product name"]).fill = PatternFill("solid", fgColor="FFF3CD")
    widths = {"Decision": 14, "Photo": 14, "Item ID": 14, "Supplier": 18, "Posted": 11, "Product name": 44, "Category": 18, "Supplier price (cost)": 12,
              "Proposed sell price": 12, "Final sell price": 12, "Sizes": 10, "Suggested": 44, "Already on site": 34, "Site price": 10,
              "Source group": 26, "Photo files": 30, "Original message": 70, "Applied": 12, "Site product ID": 38}
    for h, w in widths.items(): ws.column_dimensions[get_column_letter(COL[h])].width = w
    dv = DataValidation(type="list", formula1='"ADD,SKIP,UPDATE COST"', allow_blank=True)
    ws.add_data_validation(dv); dv.add(f"A2:A{ws.max_row + 500}")
    cats = DataValidation(type="list", formula1='"' + ",".join(CATEGORIES) + '"', allow_blank=True)
    ws.add_data_validation(cats); cats.add(f"G2:G{ws.max_row + 500}")
    ws.freeze_panes = "D2"
    n = wb.create_sheet("How to review")
    for line in ["1. In column A choose ADD or SKIP for each row (blank rows are ignored). Prices of products already on the site are changed in the admin panel, never from here.",
                 "   Orange rows (⚠) are products you already sell whose new supplier price leaves you a thin margin or a loss; blue rows (↓) leave room to cut your price.",
                 "   For those, choose UPDATE COST to record the new supplier price for the admin Inventory & pricing page. Your selling price is NOT changed: do that in admin.",
                 "2. Check the product name and category (yellow = needs a name). Branded shoes: keep the '-style' wording, no brand claims.",
                 "3. 'Final sell price' is what customers pay. It starts as supplier price + margin; change it freely. The supplier price and group are saved for the admin Inventory & pricing page.",
                 "4. Save the file, then tell Claude 'apply the approved catalogue rows'.",
                 f"Margin rule: supplier price +{CFG['markup_percent']}% rounded up to the next KSh {CFG['round_up_to']}; none for your own groups (scripts/intake/config.json)."]:
        n.append([line])
    n.column_dimensions["A"].width = 120
    wb.save(path)


def note_stock(stock, product, supplier, ts, sold_out=False, price=None):
    """Remember the newest date a supplier posted (or marked sold out) a product you sell."""
    rec = stock["products"].setdefault(product["id"], {"name": product["name"], "supplier": supplier, "last_posted": ""})
    day = ts.date().isoformat()
    rec["name"] = product["name"]
    if sold_out:
        if day >= rec.get("sold_out", ""): rec["sold_out"] = day
        return
    if day >= rec["last_posted"]:
        rec.update(last_posted=day, supplier=supplier)
        if price: rec["price"] = price
    if rec.get("sold_out") and rec["sold_out"] < rec["last_posted"]:
        rec.pop("sold_out")  # reposted after the sold-out notice: back in stock


def track_export_stock(stock, msgs, items, supplier, live):
    """Updates last-posted / sold-out dates from one export. Returns the cutoff for 'recent' posts."""
    latest = max((m_["ts"] for m_ in msgs), default=datetime.datetime.now())
    sup = stock["suppliers"].setdefault(supplier, {})
    sup["latest_chat"] = max(sup.get("latest_chat", ""), latest.date().isoformat())
    for m_ in msgs:  # text-only "sold out" notices
        t = clean_text(m_["body"])
        if not SOLD_OUT.search(t): continue
        hit = match_live(clean_name(SOLD_OUT.sub(" ", t)), live)
        if hit: note_stock(stock, hit, supplier, m_["ts"], sold_out=True)
    for it in items:
        if not it["name"]: continue
        hit = match_live(it["name"], live)
        if hit: note_stock(stock, hit, supplier, it["ts"], sold_out=bool(SOLD_OUT.search(it["text"] or "")), price=it["price"])
    return latest - datetime.timedelta(days=NEW_ITEM_DAYS)


def cmd_stock(args):
    """Stock check over every export in the Drive folder, including ones already scanned.
    Writes no review rows: only intake/stock_seen.json and the stock-check sheet."""
    root = work_root(args)
    for d in ("work", "review"): os.makedirs(os.path.join(root, d), exist_ok=True)
    watch = [cfg_path(p) for p in (args.watch or CFG["watch_dirs"])]
    zips = sorted({z for d in watch for z in glob.glob(os.path.join(d, "WhatsApp Chat*.zip"))})
    if not zips: print("No supplier chat exports found."); return
    live = live_products()
    spath = os.path.join(root, "stock_seen.json")
    stock = json.load(open(spath)) if os.path.exists(spath) else {"products": {}, "suppliers": {}}
    for z in zips:
        if not fully_local(z):
            print(f"  skipped {os.path.basename(z)}: still downloading from Google Drive (set it to 'Available offline' and wait)"); continue
        h = file_sha1(z)
        group = re.sub(r"\s*\(\d+\)$", "", re.sub(r"^WhatsApp Chat\s*-\s*", "", os.path.basename(z)[:-4])).strip()
        supplier, kind = canonical_supplier(group)
        dest = os.path.join(root, "work", f"{slug(supplier)}-{h[:8]}")
        if not glob.glob(os.path.join(dest, "*.txt")):
            shutil.rmtree(dest, ignore_errors=True); safe_extract(z, dest)
        txts = glob.glob(os.path.join(dest, "*.txt"))
        if not txts: continue
        msgs = parse_chat(open(txts[0], encoding="utf-8", errors="ignore").read())
        track_export_stock(stock, msgs, build_items(msgs, dest, supplier), supplier, live)
        print(f"  {supplier}: newest message {stock['suppliers'][supplier]['latest_chat']}")
    json.dump(stock, open(spath, "w"), indent=1)
    flagged = stock_check(stock, live)
    sold_n = sum(1 for f in flagged if f["status"].startswith("SOLD"))
    note = f"{len(zips)} export(s): {sold_n} sold out, {len(flagged) - sold_n} not reposted {STOCK_CHECK_DAYS}+ days, {len(stock['products'])} on-site product(s) tracked"
    if flagged:
        out = os.path.join(root, "review", f"{datetime.date.today().isoformat()}-stock-check.xlsx")
        write_stock_sheet(out, flagged); note += f" -> {os.path.basename(out)}"; print("wrote", out)
    print("stock check:", note)
    log_run(root, "stock", note)


def stock_check(stock, live):
    """Products still on the site that a supplier marked sold out, or hasn't reposted for a while.
    Days are counted to that supplier's newest message, so a late export doesn't make everything stale."""
    on_site = {p["id"]: p for p in live}
    out = []
    for pid, rec in stock["products"].items():
        if pid not in on_site: continue
        latest = stock["suppliers"].get(rec["supplier"], {}).get("latest_chat") or rec["last_posted"]
        if rec.get("sold_out"):
            status = f"SOLD OUT notice on {rec['sold_out']}"
        elif rec["last_posted"]:
            age = (datetime.date.fromisoformat(latest) - datetime.date.fromisoformat(rec["last_posted"])).days
            if age < STOCK_CHECK_DAYS: continue
            status = f"Not reposted for {age} days"
        else:
            continue
        out.append({"name": on_site[pid]["name"], "supplier": rec["supplier"], "last_posted": rec["last_posted"] or "-",
                    "status": status, "price": on_site[pid].get("price"), "id": pid})
    return sorted(out, key=lambda f: (not f["status"].startswith("SOLD"), f["last_posted"]))


def stale_exports(stock):
    today = datetime.date.today()
    out = []
    for name, sup in sorted(stock["suppliers"].items()):
        if not sup.get("latest_chat"): continue
        age = (today - datetime.date.fromisoformat(sup["latest_chat"])).days
        if age >= STALE_EXPORT_DAYS: out.append(f"{name} chat is {age} days old, export it again")
    return out


def write_stock_sheet(path, flagged):
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Stock check"
    head = ["Product (on site)", "Supplier", "Last posted", "Status", "Site price", "Site product ID"]
    ws.append(head)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="4D66B0")
    for f in flagged:
        ws.append([f["name"], f["supplier"], f["last_posted"], f["status"], f["price"], f["id"]])
        if f["status"].startswith("SOLD"):
            for c in ws[ws.max_row]: c.fill = PatternFill("solid", fgColor="FFE0B2")
    for col, w in zip("ABCDEF", (48, 22, 12, 34, 11, 38)): ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    n = wb.create_sheet("How to use")
    for line in [f"Products still on your site that a supplier marked sold out, or hasn't reposted for {STOCK_CHECK_DAYS}+ days.",
                 "Nothing here is changed automatically. Confirm stock with the supplier, then hide or update the product in admin if needed.",
                 f"Change the {STOCK_CHECK_DAYS}-day limit with stock_check_days in scripts/intake/config.json."]:
        n.append([line])
    n.column_dimensions["A"].width = 120
    wb.save(path)


def cmd_scan(args):
    root = work_root(args)
    for d in ("work", "review", "photos"): os.makedirs(os.path.join(root, d), exist_ok=True)
    mpath = os.path.join(root, "processed.json")
    manifest = json.load(open(mpath)) if os.path.exists(mpath) else {"zips": {}, "seen": []}
    seen = set(manifest["seen"])
    watch = [cfg_path(p) for p in (args.watch or CFG["watch_dirs"])]
    zips = sorted({z for d in watch for z in glob.glob(os.path.join(d, "WhatsApp Chat*.zip"))})
    if not zips:
        print("No supplier chat exports found in:", *watch, sep="\n  ")
        log_run(root, "scan", "No chat exports in the Drive folder"); return
    live = live_products()
    print(f"live catalogue: {len(live)} products")
    new_sheets, summary = 0, []
    cost_flags = collections.Counter()
    spath = os.path.join(root, "stock_seen.json")
    stock = json.load(open(spath)) if os.path.exists(spath) else {"products": {}, "suppliers": {}}
    new_exports = old_total = 0
    not_ready = []
    for z in zips:
        if not fully_local(z):
            print(f"  skipped {os.path.basename(z)}: still downloading from Google Drive (set it to 'Available offline' and wait)")
            not_ready.append(os.path.basename(z)); continue
        h = file_sha1(z)
        if h in manifest["zips"]:
            continue
        group = re.sub(r"\s*\(\d+\)$", "", re.sub(r"^WhatsApp Chat\s*-\s*", "", os.path.basename(z)[:-4])).strip()
        supplier, kind = canonical_supplier(group)
        dest = os.path.join(root, "work", f"{slug(supplier)}-{h[:8]}")
        shutil.rmtree(dest, ignore_errors=True); safe_extract(z, dest)
        txts = glob.glob(os.path.join(dest, "*.txt"))
        if not txts:
            print("  no chat text in", os.path.basename(z)); manifest["zips"][h] = {"zip": os.path.basename(z), "note": "no text"}; continue
        msgs = parse_chat(open(txts[0], encoding="utf-8", errors="ignore").read())
        items = build_items(msgs, dest, supplier)
        # "Recent" counts back from the newest message in this export, not from today, so a
        # chat exported late still gets a full window.
        cutoff = track_export_stock(stock, msgs, items, supplier, live)
        new_exports += 1
        old_skipped = 0
        if kind == "unlisted": print(f"  note: '{group}' is not one of your listed suppliers; using the default {CFG['markup_percent']}% margin.")
        # A whole group can be pinned to one category (config "group_categories": {"match text": "slug"}).
        group_cat = next((c for m, c in CFG.get("group_categories", {}).items() if m.lower() in group.lower()), None)
        rows, skipped_seen = [], 0
        for n, it in enumerate(items):
            sold = bool(SOLD_OUT.search(it["text"] or ""))
            if it["ts"] < cutoff: old_skipped += 1; continue  # older posts: likely gone, not offered
            if sold: continue  # a sold-out notice is never a product to add
            key = f"{slug(supplier)}|{it['ts'].date()}|{slug(it['name'] or it['photos'][0])}|{it['price']}"
            if key in seen: skipped_seen += 1; continue
            seen.add(key)
            iid = f"{slug(supplier)[:10]}-{h[:4]}-{n:03d}"
            pdir = os.path.join(root, "photos", iid); os.makedirs(pdir, exist_ok=True)
            files = []
            for i, p in enumerate(it["photos"][:CFG["photos_per_product"]], 1):
                try:
                    im = Image.open(os.path.join(dest, p)).convert("RGB"); im.thumbnail((CFG["photo_max_px"],) * 2)
                    fn = f"{iid}-{i}.jpg"; im.save(os.path.join(pdir, fn), "JPEG", quality=80, optimize=True, progressive=True); files.append(fn)
                except Exception:
                    pass
            if not files: continue
            thumb = os.path.join(pdir, "thumb.png"); t = Image.open(os.path.join(pdir, files[0])).convert("RGB"); t.thumbnail((120, 120)); t.save(thumb)
            name = it["name"] or "(name from photo)"
            m = match_live(name, live) if it["name"] else None
            cost = it["price"]; prop = sell_price(cost, kind)
            # Prices of products already on the site are never changed from here: admin sets those.
            if m is not None:
                suggested = cost_check(cost, m.get("price"), kind)
                if suggested != "SKIP (on site)": cost_flags[suggested[0]] += 1
            else: suggested = "ADD" if cost else "ADD (no price)"
            rows.append({"decision": "", "id": iid, "supplier": supplier, "posted": it["ts"].strftime("%Y-%m-%d"), "name": name,
                         "category": group_cat or it["category"], "cost": cost, "proposed": prop, "final": prop, "sizes": it["sizes"], "suggested": suggested, "group": group, "kind": kind,
                         "site_name": m["name"] if m else "", "site_price": m["price"] if m else None, "site_id": m["id"] if m else "",
                         "files": files, "text": it["text"], "thumb": thumb})
        stamp = datetime.date.today().isoformat()
        if rows:
            out = os.path.join(root, "review", f"{stamp}-{slug(supplier)}-{h[:6]}.xlsx"); write_sheet(out, rows); new_sheets += 1
            flagged = [r["id"] for r in rows if r["name"].startswith("(")]
            json.dump({"sheet": out, "photo_dir": os.path.join(root, "photos"), "need_names": flagged}, open(out.replace(".xlsx", ".todo.json"), "w"), indent=1)
            summary.append(f"{supplier} {len(rows)} new item(s) -> {os.path.basename(out)}")
            print(f"  {supplier}: {len(msgs)} messages -> {len(rows)} new items ({len(flagged)} photo-only need a name, {skipped_seen} already seen, {old_skipped} older than {NEW_ITEM_DAYS} days ignored) -> {out}")
        else:
            print(f"  {supplier}: nothing new ({old_skipped} post(s) older than {NEW_ITEM_DAYS} days ignored)")
        old_total += old_skipped
        manifest["zips"][h] = {"zip": os.path.basename(z), "supplier": supplier, "processed": stamp}
    manifest["seen"] = sorted(seen)
    json.dump(manifest, open(mpath, "w"), indent=1)
    json.dump(stock, open(spath, "w"), indent=1)
    print(f"done: {new_sheets} new review sheet(s)")
    stock_note = ""
    if new_exports:
        flagged = stock_check(stock, live)
        sold_n = sum(1 for f in flagged if f["status"].startswith("SOLD"))
        stock_note = f", {old_total} older post(s) ignored, stock check: {sold_n} sold out, {len(flagged) - sold_n} not reposted {STOCK_CHECK_DAYS}+ days"
        if flagged:
            out = os.path.join(root, "review", f"{datetime.date.today().isoformat()}-stock-check.xlsx")
            write_stock_sheet(out, flagged)
            stock_note += f" -> {os.path.basename(out)}"
            print(f"stock check: {sold_n} sold out, {len(flagged) - sold_n} not reposted for {STOCK_CHECK_DAYS}+ days -> {out}")
    stale = stale_exports(stock)
    if stale:
        print("export reminder:", "; ".join(stale)); stock_note += "; export reminder: " + "; ".join(stale)
    flags = ""
    if cost_flags:
        flags = f", {cost_flags['⚠']} thin-margin/below-cost warning(s), {cost_flags['↓']} room-to-cut"
        print(f"cost check: {cost_flags['⚠']} product(s) on site now at a thin margin or below cost, {cost_flags['↓']} with room to cut the price")
    log_run(root, "scan", f"{len(zips)} export(s) checked, {new_sheets} new review sheet(s){flags}{stock_note}" + (": " + "; ".join(summary) if summary else "") + (f" | NOT READY (still downloading from Drive): {', '.join(not_ready)}" if not_ready else ""))


def cmd_setnames(args):
    names = json.load(open(args.names))
    wb = openpyxl.load_workbook(args.sheet); ws = wb["Review"]; n = 0
    for row in ws.iter_rows(min_row=2):
        iid = row[COL["Item ID"] - 1].value
        if iid in names:
            ws.cell(row=row[0].row, column=COL["Product name"]).value = names[iid]["name"]
            if names[iid].get("category"): ws.cell(row=row[0].row, column=COL["Category"]).value = names[iid]["category"]
            ws.cell(row=row[0].row, column=COL["Product name"]).fill = PatternFill(fill_type=None); n += 1
    wb.save(args.sheet); print(f"named {n} rows")
    log_run(os.path.dirname(os.path.dirname(os.path.abspath(args.sheet))), "setnames", f"named {n} photo-only row(s) in {os.path.basename(args.sheet)}")


def cmd_contact(args):
    """Labelled contact sheets of the photo-only posts so each can be named by eye."""
    todo = json.load(open(args.todo)); ids = todo["need_names"]
    outdir = os.path.join(os.path.dirname(os.path.dirname(args.todo)), "contact"); os.makedirs(outdir, exist_ok=True)
    from PIL import ImageDraw, ImageFont
    try: font = ImageFont.truetype("/System/Library/Fonts/Supplemental/Arial Bold.ttf", 34)
    except Exception: font = ImageFont.load_default()
    per, cell, made = 12, 380, []
    for s0 in range(0, len(ids), per):
        chunk = ids[s0:s0 + per]
        sheet = Image.new("RGB", (4 * cell, ((len(chunk) + 3) // 4) * cell), "white"); d = ImageDraw.Draw(sheet)
        for k, iid in enumerate(chunk):
            fn = sorted(glob.glob(os.path.join(todo["photo_dir"], iid, f"{iid}-1.jpg")))
            if not fn: continue
            im = Image.open(fn[0]).convert("RGB"); im.thumbnail((cell - 8, cell - 8))
            x, y = (k % 4) * cell, (k // 4) * cell; sheet.paste(im, (x + 4, y + 4))
            d.rectangle([x + 4, y + 4, x + 96, y + 46], fill="black"); d.text((x + 10, y + 6), iid[-3:], font=font, fill="yellow")
        path = os.path.join(outdir, f"{os.path.basename(args.todo)[:-10]}-{s0 // per}.jpg"); sheet.save(path, quality=80); made.append(path)
    print("\n".join(made)); print("label = last 3 digits of the Item ID; for the names file use the full Item ID")


# ---------------------------------------------------------------- apply
def sourcing_ready(base, H):
    st, body = http("GET", f"{base}/rest/v1/product_sourcing?select=product_id&limit=1", H)
    return st == 200


SQL_HINT = "The admin-only table product_sourcing doesn't exist yet. Run docs/sql/product_sourcing.sql in Supabase (SQL Editor) once, then retry."


def update_cost_row(row, base, H, dry_run):
    """Records a supplier's new price as the product's cost in product_sourcing.
    Never touches the product's selling price. Returns True (done), False (needs attention)."""
    v = lambda h: row[COL[h] - 1].value if len(row) >= COL[h] else None
    pid, cost, site = v("Site product ID"), v("Supplier price (cost)"), v("Already on site") or v("Product name")
    if not pid: print(f"  ! {v('Item ID')}: UPDATE COST needs a product already on the site (sheet made before this feature: re-scan)"); return False
    if not isinstance(cost, (int, float)) or cost <= 0: print(f"  ! {site}: UPDATE COST needs a supplier price"); return False
    supplier = canonical_supplier(v("Source group") or "")[0] if v("Source group") else (v("Supplier") or "Unknown")
    print(f"  {'(dry) ' if dry_run else ''}cost  {site} | supplier cost now KSh {int(cost)} from {supplier} (selling price unchanged)")
    if dry_run: return True
    patch = {"supplier_cost": int(cost), "supplier": supplier, "source_ref": f"cost update: {v('Source group')} post of {v('Posted')}"}
    hp = dict(H, Prefer="return=representation")
    st, body = http("PATCH", f"{base}/rest/v1/product_sourcing?product_id=eq.{pid}", hp, json.dumps(patch).encode())
    if st not in (200, 204): print(f"    failed: {st} {body[:120]}"); return False
    if st == 200 and json.loads(body or b"[]"): return True
    # No sourcing record yet: create one, listed at today's selling price.
    st, body = http("GET", f"{base}/rest/v1/products?select=price&id=eq.{pid}", H)
    rows = json.loads(body) if st == 200 else []
    if not rows: print(f"    failed: product {pid} not found"); return False
    rec = dict(patch, product_id=pid, listed_price=rows[0]["price"])
    st, body = http("POST", f"{base}/rest/v1/product_sourcing", H, json.dumps(rec).encode())
    if st not in (200, 201): print(f"    failed: {st} {body[:120]}"); return False
    return True


def cmd_apply(args):
    root = work_root(args)
    key = os.environ.get("SUPABASE_SERVICE_KEY")
    if not key and not args.dry_run:
        sys.exit("SUPABASE_SERVICE_KEY is not set (pass it inline for this run only).")
    env = load_env(); base = env["NEXT_PUBLIC_SUPABASE_URL"]
    H = {"apikey": key or "", "Authorization": f"Bearer {key}", "Content-Type": "application/json"}
    HR = H if key else {"apikey": env["NEXT_PUBLIC_SUPABASE_ANON_KEY"]}
    sheets = [args.sheet] if args.sheet else sorted(f for f in glob.glob(os.path.join(root, "review", "*.xlsx")) if not os.path.basename(f).startswith("~$"))
    if key and not sourcing_ready(base, H): sys.exit(SQL_HINT)
    st, body = http("GET", f"{base}/rest/v1/products?select=slug", HR)
    taken = {r["slug"] for r in json.loads(body)}
    st, body = http("GET", f"{base}/rest/v1/categories?select=slug", HR)
    cats = {r["slug"] for r in json.loads(body)}
    added = bad = costs = 0
    for sp in sheets:
        wb = openpyxl.load_workbook(sp)
        if "Review" not in wb.sheetnames: continue  # e.g. the stock-check sheet
        ws = wb["Review"]; dirty = False
        for row in ws.iter_rows(min_row=2):
            v = lambda h: row[COL[h] - 1].value
            decision = (v("Decision") or "").strip().upper()
            if v("Applied"): continue
            if decision == "UPDATE COST":
                r_ = update_cost_row(row, base, H, args.dry_run)
                if r_ is True:
                    costs += 1
                    if not args.dry_run:
                        ws.cell(row=row[0].row, column=COL["Applied"]).value = "cost " + datetime.date.today().isoformat(); dirty = True
                elif r_ is False: bad += 1
                continue
            if decision != "ADD": continue
            name, cat, price = (v("Product name") or "").strip(), (v("Category") or "").strip(), v("Final sell price")
            if not name or name.startswith("("): print(f"  ! {v('Item ID')}: needs a product name"); bad += 1; continue
            if cat == "sneakers-shoes" and BIG_BRANDS.search(name) and "style" not in name.lower():
                print(f"  ! {name}: rename with neutral '-style' wording (e.g. 'Cushion Sneakers (Air Max style)'), no brand claims"); bad += 1; continue
            if not isinstance(price, (int, float)) or price < CFG["min_sell_price"]: print(f"  ! {name}: needs a final sell price"); bad += 1; continue
            if cat not in cats: print(f"  ! {name}: category '{cat}' doesn't exist yet (create it first)"); bad += 1; continue
            kind = canonical_supplier(v("Source group") or "")[1]
            cost = v("Supplier price (cost)") if kind != "retail" and isinstance(v("Supplier price (cost)"), (int, float)) else None
            supplier = canonical_supplier(v("Source group") or "")[0] if v("Source group") else (v("Supplier") or "Unknown")
            base_slug = s_ = slug(name); n = 2
            while s_ in taken: s_ = f"{base_slug}-{n}"; n += 1
            taken.add(s_)
            files = [f for f in (v("Photo files") or "").split(", ") if f]
            pdir = os.path.join(root, "photos", v("Item ID"))
            sizes = v("Sizes")
            desc = DESC.get(cat, "") + (f" Sizes {sizes}. Please add your size in the Order Notes at checkout." if sizes and cat == "sneakers-shoes" else "")
            print(f"  {'(dry) ' if args.dry_run else ''}add   {name} | {cat} | KSh {int(price)} (cost {cost or 'n/a'} from {supplier}) | {len(files)} photo(s)")
            if args.dry_run: added += 1; continue
            pid = str(uuid.uuid4()); urls = []
            try:
                for fn in files:
                    hd = {"apikey": key, "Authorization": f"Bearer {key}", "Content-Type": "image/jpeg", "x-upsert": "true"}
                    st_, r = http("POST", f"{base}/storage/v1/object/product-images/{pid}/{fn}", hd, open(os.path.join(pdir, fn), "rb").read())
                    if st_ not in (200, 201): raise RuntimeError(f"photo {fn}: {st_}")
                    urls.append(f"{base}/storage/v1/object/public/product-images/{pid}/{fn}")
                rowj = {"id": pid, "slug": s_, "name": name, "category": cat, "subcategory": None, "price": int(price), "discount_price": None,
                        "images": urls, "colors": None, "features": [], "rating": 0, "review_count": 0, "stock": 0, "is_deal": False,
                        "deal_ends_at": None, "deal_claimed_percent": None, "is_new_arrival": True, "verified": False, "description": desc, "jumia_url": None}
                st_, r = http("POST", f"{base}/rest/v1/products", H, json.dumps(rowj).encode())
                if st_ not in (200, 201): raise RuntimeError(f"insert: {st_} {r[:150]}")
                src = {"product_id": pid, "supplier": supplier, "supplier_cost": int(cost) if cost else None, "listed_price": int(price),
                       "source_ref": f"{v('Source group')} export, post of {v('Posted')}"}
                st_, r = http("POST", f"{base}/rest/v1/product_sourcing", H, json.dumps(src).encode())
                if st_ not in (200, 201): print(f"    warning: product added but its sourcing record failed ({st_} {r[:120]})")
                added += 1
                ws.cell(row=row[0].row, column=COL["Applied"]).value = datetime.date.today().isoformat(); dirty = True
            except Exception as e:
                print("    failed:", e)
        if dirty: wb.save(sp)
    print(f"{'(dry run) ' if args.dry_run else ''}added {added}, supplier costs updated {costs}, rows needing attention {bad}")
    log_run(root, "apply", f"{'(dry run) ' if args.dry_run else ''}added {added} product(s), updated {costs} supplier cost(s), {bad} row(s) need attention, from {len(sheets)} sheet(s)")
    if added and not args.dry_run:
        print("Catalogue changed in the database. It shows on the site after the next rebuild (Publish changes in admin).")


def cmd_backfill(args):
    """One-off: give products that have no sourcing record one (supplier unknown / own group), so the
    admin Inventory & pricing page covers the whole catalogue. Costs stay blank where unknown."""
    key = os.environ.get("SUPABASE_SERVICE_KEY") or sys.exit("SUPABASE_SERVICE_KEY is not set (pass it inline for this run only).")
    env = load_env(); base = env["NEXT_PUBLIC_SUPABASE_URL"]
    H = {"apikey": key, "Authorization": f"Bearer {key}", "Content-Type": "application/json"}
    if not sourcing_ready(base, H): sys.exit(SQL_HINT)
    own = set(json.load(open(args.own_slugs))) if args.own_slugs else set()
    st, body = http("GET", f"{base}/rest/v1/products?select=id,slug,price,created_at", H); prods = json.loads(body)
    st, body = http("GET", f"{base}/rest/v1/product_sourcing?select=product_id", H); have = {r["product_id"] for r in json.loads(body)}
    rows = []
    for p in prods:
        if p["id"] in have: continue
        is_own = p["slug"] in own
        rows.append({"product_id": p["id"], "supplier": CFG["own_group_label"] if is_own else "Earlier import (supplier not recorded)",
                     "supplier_cost": None, "listed_price": p["price"], "source_ref": "backfilled " + datetime.date.today().isoformat(),
                     "listed_at": p.get("created_at") or datetime.datetime.utcnow().isoformat()})
    for i in range(0, len(rows), 100):
        st, r = http("POST", f"{base}/rest/v1/product_sourcing", H, json.dumps(rows[i:i + 100]).encode())
        if st not in (200, 201): sys.exit(f"backfill failed ({st}): {r[:200]}")
    print(f"backfilled {len(rows)} products ({sum(1 for r in rows if r['supplier'] == CFG['own_group_label'])} from the offers group)")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    s = sub.add_parser("scan"); s.add_argument("--watch", action="append"); s.add_argument("--work")
    n = sub.add_parser("setnames"); n.add_argument("sheet"); n.add_argument("names")
    c = sub.add_parser("contact"); c.add_argument("todo")
    b = sub.add_parser("backfill"); b.add_argument("--own-slugs")
    k = sub.add_parser("stock"); k.add_argument("--watch", action="append"); k.add_argument("--work")
    g = sub.add_parser("log"); g.add_argument("--last", type=int, default=15); g.add_argument("--work")
    a = sub.add_parser("apply"); a.add_argument("--sheet"); a.add_argument("--work"); a.add_argument("--dry-run", action="store_true")
    args = ap.parse_args()
    {"scan": cmd_scan, "setnames": cmd_setnames, "contact": cmd_contact, "apply": cmd_apply, "backfill": cmd_backfill, "log": cmd_log, "stock": cmd_stock}[args.cmd](args)


if __name__ == "__main__":
    main()
