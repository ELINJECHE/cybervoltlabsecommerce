#!/usr/bin/env python3
"""Build exports/cutout_review.xlsx from intake/cutouts/results.json: before/after thumbnails, a suggested YES/NO,
and the reason. The owner edits column A, then `apply_cutouts.py` uploads only the YES rows (keeping originals)."""
import os, json, io, urllib.request, concurrent.futures as cf
import openpyxl
from openpyxl.drawing.image import Image as XImage
from openpyxl.styles import Font, PatternFill
from PIL import Image
ROOT = os.path.join(os.path.dirname(__file__), "..", "..")
OUT = os.path.join(ROOT, "intake", "cutouts"); res = json.load(open(os.path.join(OUT, "results.json")))
env = {}
for l in open(os.path.join(ROOT, ".env.local")):
    if "=" in l and not l.startswith("#"): a, b = l.strip().split("=", 1); env[a] = b
u, k = env["NEXT_PUBLIC_SUPABASE_URL"], env["NEXT_PUBLIC_SUPABASE_ANON_KEY"]
prods, off = [], 0
while True:
    r = urllib.request.Request(f"{u}/rest/v1/products?select=slug,name,category,images&order=category,name", headers={"apikey": k, "Authorization": "Bearer " + k, "Range": f"{off}-{off+999}"})
    d = json.load(urllib.request.urlopen(r)); prods += d
    if len(d) < 1000: break
    off += 1000
SKIP_CATS = {"interior-design": "bedding and fabrics cut out badly"}
def verdict(r, cat):
    if r["status"] != "ok": return "NO", r.get("why", r["status"])
    if min(r["w"], r["h"]) < 400: return "NO", "original photo too small"
    if r["edges"] >= 3: return "NO", "subject fills the photo (poster / pack shot / cropped)"
    if r["frac"] > 0.88: return "NO", "subject fills the photo"
    if r["frac"] < 0.10: return "NO", "subject very small or not found"
    if cat in SKIP_CATS: return "NO", SKIP_CATS[cat]
    return "YES", "clean cut-out expected"
def thumb(data, px=110):
    im = Image.open(io.BytesIO(data)).convert("RGB"); im.thumbnail((px, px)); b = io.BytesIO(); im.save(b, "JPEG", quality=70); b.seek(0); return b
def before(p):
    try:
        url = p["images"][0]; url = url if url.startswith("http") else "https://cybervoltlabs.co.ke" + url
        return p["slug"], thumb(urllib.request.urlopen(url, timeout=40).read())
    except Exception: return p["slug"], None
rows = [p for p in prods if p["slug"] in res and (os.path.exists(os.path.join(OUT, p["slug"] + ".jpg")) or res[p["slug"]]["status"] != "ok")]
with cf.ThreadPoolExecutor(4) as ex: bef = dict(ex.map(before, rows))
wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Cut-outs"
hdr = ["Use cut-out? (YES/NO)", "Before", "After", "Product", "Category", "Why", "Slug"]; ws.append(hdr)
for c in ws[1]: c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="305496")
n = 0; yes = 0
for p in rows:
    r = res[p["slug"]]; v, why = verdict(r, p["category"]); ws.append([v, None, None, p["name"], p["category"], why, p["slug"]]); i = ws.max_row; ws.row_dimensions[i].height = 85
    yes += v == "YES"
    if bef.get(p["slug"]): ws.add_image(XImage(bef[p["slug"]]), f"B{i}")
    f = os.path.join(OUT, p["slug"] + ".jpg")
    if os.path.exists(f): ws.add_image(XImage(thumb(open(f, "rb").read())), f"C{i}")
for col, w in zip("ABCDEFG", [16, 17, 17, 48, 22, 44, 34]): ws.column_dimensions[col].width = w
ws.freeze_panes = "A2"
os.makedirs(os.path.join(ROOT, "exports"), exist_ok=True)
dest = os.path.join(ROOT, "exports", "cutout_review.xlsx"); wb.save(dest)
print("rows:", len(rows), "| suggested YES:", yes, "| saved", dest)
